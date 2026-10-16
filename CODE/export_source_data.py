from pathlib import Path
import json,csv
from openpyxl import Workbook,load_workbook
R=Path(__file__).resolve().parents[1];out=R/'REPRODUCED/source_data';out.mkdir(parents=True,exist_ok=True)
payload=json.loads((R/'SOURCE_DATA/frozen_source_data_payload.json').read_text())
w=Workbook();w.remove(w.active)
original=load_workbook(R/'SOURCE_DATA/Source_Data_V2.xlsx',read_only=True,data_only=True)
checked=0
for obj in payload:
 sheet=w.create_sheet(obj['name'])
 for row in obj['matrix']:sheet.append(row)
 with (out/(obj['name']+'.tsv')).open('w',newline='') as f:csv.writer(f,delimiter='\t').writerows(obj['matrix'])
 expected=list(original[obj['name']].values)
 for i,row in enumerate(obj['matrix']):
  for j,value in enumerate(row):
   assert value==expected[i][j],(obj['name'],i,j,value,expected[i][j])
   checked+=1
w.save(out/'Source_Data.xlsx')
print('PASS frozen workbook cells:',checked)
